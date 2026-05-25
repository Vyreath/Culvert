import os, csv, io
from datetime import datetime
from pathlib import Path
from io import BytesIO

from dotenv import load_dotenv
from flask import Flask, jsonify, request, send_from_directory, Response, g
from flask_cors import CORS
from supabase import create_client
from functools import wraps

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                Paragraph, Spacer, PageBreak, KeepTogether)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER, TA_LEFT

BASE_DIR = Path(__file__).resolve().parent
ENV_PATH = BASE_DIR / ".env"
load_dotenv(dotenv_path=ENV_PATH)

# ═════════════════════════════════════════════
# CONEXIÓN SUPABASE
# ═════════════════════════════════════════════
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY") or os.getenv("SUPABASE_SERVICE_ROLE_KEY")
if not SUPABASE_URL or not SUPABASE_KEY:
    raise RuntimeError(f"Faltan credenciales de Supabase. Configura SUPABASE_URL y SUPABASE_KEY en {ENV_PATH}")

supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
supabase_admin = create_client(SUPABASE_URL, SUPABASE_KEY)

FRONT_DIR = Path(os.getenv("FRONT_DIR", BASE_DIR / "front")).resolve()

app = Flask(__name__)
CORS(app)

# ─────────────────────────────────────────────
# DECORADOR DE AUTENTICACIÓN
# ─────────────────────────────────────────────
def require_auth(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        auth_header = request.headers.get("Authorization", "")
        token = auth_header.removeprefix("Bearer ").strip()
        if not token:
            return jsonify({"error": "Token requerido"}), 401
        try:
            user_resp = supabase_admin.auth.get_user(token)
            if not user_resp or not user_resp.user:
                raise Exception("Token inválido")
            g.current_user = user_resp.user
        except Exception:
            return jsonify({"error": "Token inválido o expirado"}), 401
        return f(*args, **kwargs)
    return decorated

def current_user_role():
    user = g.get("current_user")
    if not user:
        return None
    metadata = user.user_metadata or {}
    return metadata.get("role", "inspector")

# ─────────────────────────────────────────────
# FRONTEND (SPA)
# ─────────────────────────────────────────────
@app.route("/")
def index():
    return send_from_directory(FRONT_DIR, "index.html")

@app.route("/login")
def login_page():
    return send_from_directory(FRONT_DIR, "login.html")

@app.route("/tuberias")
def tuberias_page():
    return send_from_directory(FRONT_DIR, "tuberias.html")

@app.route("/muros")
def muros_page():
    return send_from_directory(FRONT_DIR, "muros.html")

@app.route("/pozos")
def pozos_page():
    return send_from_directory(FRONT_DIR, "pozos.html")

@app.route("/inspecciones")
def inspecciones_page():
    return send_from_directory(FRONT_DIR, "inspecciones.html")

@app.route("/mapa")
def mapa_page():
    return send_from_directory(FRONT_DIR, "mapa.html")

@app.route("/auditoria")
def auditoria_page():
    return send_from_directory(FRONT_DIR, "auditoria.html")

@app.route("/css/<path:filename>")
def css_files(filename):
    return send_from_directory(FRONT_DIR / "css", filename)

@app.route("/js/<path:filename>")
def js_files(filename):
    return send_from_directory(FRONT_DIR / "js", filename)

# ═════════════════════════════════════════════
# AUTENTICACIÓN
# ═════════════════════════════════════════════
DEFAULT_PERMISSIONS = {
    "alcantarillas_crear": True,
    "alcantarillas_editar": True,
    "alcantarillas_eliminar": False,
    "catalogos_editar": False,
}

def build_auth_payload(auth_response):
    session = getattr(auth_response, "session", None)
    user = getattr(auth_response, "user", None)
    if not user:
        return None
    metadata = getattr(user, "user_metadata", None) or {}
    app_metadata = getattr(user, "app_metadata", None) or {}
    role = app_metadata.get("role") or metadata.get("role") or "inspector"
    payload = {
        "user": {
            "id": getattr(user, "id", None),
            "email": getattr(user, "email", None),
            "nombre": metadata.get("full_name") or metadata.get("nombre") or getattr(user, "email", None),
            "role": role,
            "permisos": metadata.get("permisos") or DEFAULT_PERMISSIONS,
        },
        "requires_email_confirmation": session is None,
    }
    if session:
        payload["access_token"] = getattr(session, "access_token", None)
        payload["refresh_token"] = getattr(session, "refresh_token", None)
    return payload

@app.route("/api/auth/register", methods=["POST"])
def register():
    body = request.json or {}
    email = (body.get("email") or "").strip().lower()
    password = body.get("password") or ""
    nombre = (body.get("nombre") or body.get("full_name") or "").strip()
    if not email or not password:
        return jsonify({"detail": "Correo y contraseña son obligatorios"}), 400
    if len(password) < 6:
        return jsonify({"detail": "La contraseña debe tener al menos 6 caracteres"}), 400
    try:
        auth_response = supabase_admin.auth.sign_up({
            "email": email,
            "password": password,
            "options": {
                "data": {
                    "full_name": nombre,
                    "role": "inspector",
                    "permisos": DEFAULT_PERMISSIONS,
                }
            },
        })
    except Exception as exc:
        return jsonify({"detail": str(exc)}), 400
    payload = build_auth_payload(auth_response)
    if not payload:
        return jsonify({"detail": "No se pudo crear el usuario"}), 400
    audit_log("register", f"Usuario {email} registrado")
    return jsonify(payload), 201

@app.route("/api/auth/login", methods=["POST"])
def login():
    body = request.json or {}
    email = (body.get("email") or "").strip().lower()
    password = body.get("password") or ""
    if not email or not password:
        return jsonify({"detail": "Correo y contraseña son obligatorios"}), 400
    try:
        auth_response = supabase_admin.auth.sign_in_with_password({
            "email": email,
            "password": password,
        })
    except Exception:
        return jsonify({"error": "Correo o contraseña incorrectos"}), 401
    payload = build_auth_payload(auth_response)
    if not payload or not payload.get("access_token"):
        return jsonify({"error": "No se pudo iniciar sesión"}), 401
    audit_log("login", f"Usuario {email} inició sesión")
    return jsonify(payload)

@app.route("/api/auth/refresh", methods=["POST"])
def refresh():
    auth_header = request.headers.get("Authorization", "")
    refresh_token = auth_header.removeprefix("Bearer ").strip()
    if not refresh_token:
        return jsonify({"error": "Token de actualización requerido"}), 401
    try:
        auth_response = supabase_admin.auth.refresh_session(refresh_token)
    except Exception:
        return jsonify({"error": "Sesión expirada"}), 401
    payload = build_auth_payload(auth_response)
    if not payload or not payload.get("access_token"):
        return jsonify({"error": "No se pudo renovar la sesión"}), 401
    return jsonify(payload)

@app.route("/api/auth/logout", methods=["POST"])
@require_auth
def logout():
    try:
        supabase_admin.auth.sign_out()
        audit_log("logout", f"Usuario {g.current_user.email} cerró sesión")
    except:
        pass
    return jsonify({"message": "Sesión cerrada"})

# ═════════════════════════════════════════════
# AUDITORÍA
# ═════════════════════════════════════════════
def audit_log(accion, detalle):
    try:
        supabase.table("auditoria").insert({
            "usuario": getattr(g.get("current_user"), "email", "sistema"),
            "accion": accion,
            "detalle": detalle,
            "fecha": datetime.utcnow().isoformat()
        }).execute()
    except:
        pass

@app.route("/api/audit", methods=["GET"])
@require_auth
def get_audit():
    if current_user_role() != "administrador":
        return jsonify({"error": "No autorizado"}), 403
    page = int(request.args.get("page", 1))
    per_page = min(int(request.args.get("per_page", 20)), 100)
    offset = (page - 1) * per_page
    data = supabase.table("auditoria").select("*").order("fecha", desc=True).range(offset, offset + per_page - 1).execute()
    count = supabase.table("auditoria").select("id", count="exact").execute().count or 0
    return jsonify({
        "data": data.data,
        "total": count,
        "page": page,
        "per_page": per_page,
        "pages": max(1, -(-count // per_page))
    })

@app.route("/api/audit/stats", methods=["GET"])
@require_auth
def audit_stats():
    if current_user_role() != "administrador":
        return jsonify({"error": "No autorizado"}), 403
    total = supabase.table("auditoria").select("id", count="exact").execute().count or 0
    return jsonify({"total": total})

# ═════════════════════════════════════════════
# CATÁLOGOS
# ═════════════════════════════════════════════
@app.route("/api/estados", methods=["GET"])
@require_auth
def get_estados():
    data = supabase.table("estados").select("*").order("id").execute()
    return jsonify(data.data)

@app.route("/api/materiales", methods=["GET"])
@require_auth
def get_materiales():
    data = supabase.table("materiales").select("*").order("id").execute()
    return jsonify(data.data)

# ═════════════════════════════════════════════
# ALCANTARILLAS
# ═════════════════════════════════════════════
@app.route("/api/alcantarillas", methods=["GET"])
@require_auth
def get_alcantarillas():
    query = supabase.table("alcantarillas").select("*")
    provincia = request.args.get("provincia")
    canton = request.args.get("canton")
    universidad = request.args.get("universidad")
    q = (request.args.get("q") or "").strip()
    if provincia: query = query.ilike("provincia", f"%{provincia}%")
    if canton: query = query.ilike("canton", f"%{canton}%")
    if universidad: query = query.ilike("universidad", f"%{universidad}%")
    if q:
        term = f"%{q}%"
        q_filters = [
            f"ubicacion.ilike.{term}", f"parroquia.ilike.{term}", f"canton.ilike.{term}",
            f"provincia.ilike.{term}", f"universidad.ilike.{term}", f"tramo_vial.ilike.{term}",
        ]
        if q.isdigit(): q_filters.insert(0, f"ficha_numero.eq.{q}")
        query = query.or_(",".join(q_filters))
    page = int(request.args.get("page", 1))
    per_page = min(int(request.args.get("per_page", 50)), 500)
    offset = (page - 1) * per_page
    query = query.order("ficha_numero").range(offset, offset + per_page - 1)
    data = query.execute()
    count_q = supabase.table("alcantarillas").select("id", count="exact")
    if provincia: count_q = count_q.ilike("provincia", f"%{provincia}%")
    if canton: count_q = count_q.ilike("canton", f"%{canton}%")
    if universidad: count_q = count_q.ilike("universidad", f"%{universidad}%")
    if q:
        if q.isdigit(): count_q = count_q.or_(f"ficha_numero.eq.{q}")
        else: count_q = count_q.or_(",".join(q_filters))
    total = count_q.execute().count or 0
    return jsonify({
        "data": data.data,
        "total": total,
        "page": page,
        "per_page": per_page,
        "pages": max(1, -(-total // per_page)),
    })

@app.route("/api/alcantarillas/<int:id>", methods=["GET"])
@require_auth
def get_alcantarilla(id):
    alc = supabase.table("alcantarillas").select("*").eq("id", id).execute()
    if not alc.data: return jsonify({"detail": "No encontrada"}), 404
    row = alc.data[0]
    row["tuberias"]   = supabase.table("tuberias").select("*, materiales(nombre), estados(nombre)").eq("alcantarilla_id", id).execute().data
    row["muros"]      = supabase.table("muros").select("*, materiales(nombre), estados(nombre)").eq("alcantarilla_id", id).execute().data
    row["pozos"]      = supabase.table("pozos_recoleccion").select("*, estados(nombre)").eq("alcantarilla_id", id).execute().data
    row["fotografias"] = supabase.table("fotografias").select("*").eq("alcantarilla_id", id).execute().data
    row["inspecciones"] = supabase.table("inspecciones").select("*").eq("alcantarilla_id", id).order("fecha", desc=True).execute().data
    row["archivos_adjuntos"] = supabase.table("archivos_adjuntos").select("*").eq("alcantarilla_id", id).execute().data
    return jsonify(row)

@app.route("/api/alcantarillas", methods=["POST"])
@require_auth
def create_alcantarilla():
    body = request.json or {}
    required = ["ficha_numero", "ubicacion"]
    if not all(k in body and body[k] for k in required):
        return jsonify({"detail": f"Campos obligatorios: {required}"}), 400
    payload = {k: v for k, v in body.items() if v not in (None, "", [])}
    result = supabase.table("alcantarillas").insert(payload).execute()
    audit_log("crear_alcantarilla", f"Ficha {result.data[0].get('ficha_numero')} creada")
    return jsonify(result.data[0]), 201

@app.route("/api/alcantarillas/<int:id>", methods=["PUT"])
@require_auth
def update_alcantarilla(id):
    body = request.json or {}
    if not body: return jsonify({"detail": "Sin datos"}), 400
    result = supabase.table("alcantarillas").update(body).eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrada"}), 404
    audit_log("actualizar_alcantarilla", f"Ficha {result.data[0].get('ficha_numero')} actualizada")
    return jsonify(result.data[0])

@app.route("/api/alcantarillas/<int:id>", methods=["DELETE"])
@require_auth
def delete_alcantarilla(id):
    result = supabase.table("alcantarillas").delete().eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrada"}), 404
    audit_log("eliminar_alcantarilla", f"Alcantarilla {id} eliminada")
    return jsonify({"message": "Eliminada"}), 200

# ═════════════════════════════════════════════
# TUBERÍAS
# ═════════════════════════════════════════════
@app.route("/api/tuberias", methods=["GET"])
@require_auth
def get_tuberias():
    alc_id = request.args.get("alcantarilla_id")
    query = supabase.table("tuberias").select("*, materiales(nombre), estados(nombre)")
    if alc_id: query = query.eq("alcantarilla_id", alc_id)
    return jsonify(query.execute().data)

@app.route("/api/tuberias/<int:id>", methods=["GET"])
@require_auth
def get_tuberia(id):
    data = supabase.table("tuberias").select("*, materiales(nombre), estados(nombre)").eq("id", id).execute()
    if not data.data: return jsonify({"detail": "No encontrada"}), 404
    return jsonify(data.data[0])

@app.route("/api/tuberias", methods=["POST"])
@require_auth
def create_tuberia():
    body = request.json or {}
    if "alcantarilla_id" not in body: return jsonify({"detail": "alcantarilla_id obligatorio"}), 400
    result = supabase.table("tuberias").insert(body).execute()
    return jsonify(result.data[0]), 201

@app.route("/api/tuberias/<int:id>", methods=["PUT"])
@require_auth
def update_tuberia(id):
    body = request.json or {}
    result = supabase.table("tuberias").update(body).eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrada"}), 404
    return jsonify(result.data[0])

@app.route("/api/tuberias/<int:id>", methods=["DELETE"])
@require_auth
def delete_tuberia(id):
    result = supabase.table("tuberias").delete().eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrada"}), 404
    return jsonify({"message": "Eliminada"}), 200

# MUROS
@app.route("/api/muros", methods=["GET"])
@require_auth
def get_muros():
    alc_id = request.args.get("alcantarilla_id")
    query = supabase.table("muros").select("*, materiales(nombre), estados(nombre)")
    if alc_id: query = query.eq("alcantarilla_id", alc_id)
    return jsonify(query.execute().data)

@app.route("/api/muros/<int:id>", methods=["GET"])
@require_auth
def get_muro(id):
    data = supabase.table("muros").select("*, materiales(nombre), estados(nombre)").eq("id", id).execute()
    if not data.data: return jsonify({"detail": "No encontrado"}), 404
    return jsonify(data.data[0])

@app.route("/api/muros", methods=["POST"])
@require_auth
def create_muro():
    body = request.json or {}
    if "alcantarilla_id" not in body: return jsonify({"detail": "alcantarilla_id obligatorio"}), 400
    result = supabase.table("muros").insert(body).execute()
    return jsonify(result.data[0]), 201

@app.route("/api/muros/<int:id>", methods=["PUT"])
@require_auth
def update_muro(id):
    body = request.json or {}
    result = supabase.table("muros").update(body).eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrado"}), 404
    return jsonify(result.data[0])

@app.route("/api/muros/<int:id>", methods=["DELETE"])
@require_auth
def delete_muro(id):
    result = supabase.table("muros").delete().eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrado"}), 404
    return jsonify({"message": "Eliminado"}), 200

# POZOS
@app.route("/api/pozos_recoleccion", methods=["GET"])
@require_auth
def get_pozos():
    alc_id = request.args.get("alcantarilla_id")
    query = supabase.table("pozos_recoleccion").select("*, estados(nombre)")
    if alc_id: query = query.eq("alcantarilla_id", alc_id)
    return jsonify(query.execute().data)

@app.route("/api/pozos_recoleccion/<int:id>", methods=["GET"])
@require_auth
def get_pozo(id):
    data = supabase.table("pozos_recoleccion").select("*, estados(nombre)").eq("id", id).execute()
    if not data.data: return jsonify({"detail": "No encontrado"}), 404
    return jsonify(data.data[0])

@app.route("/api/pozos_recoleccion", methods=["POST"])
@require_auth
def create_pozo():
    body = request.json or {}
    if "alcantarilla_id" not in body: return jsonify({"detail": "alcantarilla_id obligatorio"}), 400
    result = supabase.table("pozos_recoleccion").insert(body).execute()
    return jsonify(result.data[0]), 201

@app.route("/api/pozos_recoleccion/<int:id>", methods=["PUT"])
@require_auth
def update_pozo(id):
    body = request.json or {}
    result = supabase.table("pozos_recoleccion").update(body).eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrado"}), 404
    return jsonify(result.data[0])

@app.route("/api/pozos_recoleccion/<int:id>", methods=["DELETE"])
@require_auth
def delete_pozo(id):
    result = supabase.table("pozos_recoleccion").delete().eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrado"}), 404
    return jsonify({"message": "Eliminado"}), 200

# FOTOGRAFÍAS
@app.route("/api/fotografias", methods=["GET"])
@require_auth
def get_fotografias():
    alc_id = request.args.get("alcantarilla_id")
    query = supabase.table("fotografias").select("*")
    if alc_id: query = query.eq("alcantarilla_id", alc_id)
    return jsonify(query.execute().data)

@app.route("/api/fotografias/<int:id>", methods=["GET"])
@require_auth
def get_fotografia(id):
    data = supabase.table("fotografias").select("*").eq("id", id).execute()
    if not data.data: return jsonify({"detail": "No encontrada"}), 404
    return jsonify(data.data[0])

@app.route("/api/fotografias", methods=["POST"])
@require_auth
def create_fotografia():
    body = request.json or {}
    required = ["alcantarilla_id", "image_url"]
    if not all(k in body for k in required): return jsonify({"detail": f"Campos obligatorios: {required}"}), 400
    result = supabase.table("fotografias").insert(body).execute()
    return jsonify(result.data[0]), 201

@app.route("/api/fotografias/<int:id>", methods=["DELETE"])
@require_auth
def delete_fotografia(id):
    result = supabase.table("fotografias").delete().eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrada"}), 404
    return jsonify({"message": "Eliminada"}), 200

# INSPECCIONES
@app.route("/api/inspecciones", methods=["GET"])
@require_auth
def get_inspecciones():
    alc_id = request.args.get("alcantarilla_id")
    query = supabase.table("inspecciones").select("*").order("fecha", desc=True)
    if alc_id: query = query.eq("alcantarilla_id", alc_id)
    return jsonify(query.execute().data)

@app.route("/api/inspecciones/<int:id>", methods=["GET"])
@require_auth
def get_inspeccion(id):
    data = supabase.table("inspecciones").select("*").eq("id", id).execute()
    if not data.data: return jsonify({"detail": "No encontrada"}), 404
    return jsonify(data.data[0])

@app.route("/api/inspecciones", methods=["POST"])
@require_auth
def create_inspeccion():
    body = request.json or {}
    required = ["alcantarilla_id", "fecha"]
    if not all(k in body for k in required): return jsonify({"detail": f"Campos obligatorios: {required}"}), 400
    result = supabase.table("inspecciones").insert(body).execute()
    return jsonify(result.data[0]), 201

@app.route("/api/inspecciones/<int:id>", methods=["PUT"])
@require_auth
def update_inspeccion(id):
    body = request.json or {}
    result = supabase.table("inspecciones").update(body).eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrada"}), 404
    return jsonify(result.data[0])

@app.route("/api/inspecciones/<int:id>", methods=["DELETE"])
@require_auth
def delete_inspeccion(id):
    result = supabase.table("inspecciones").delete().eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrada"}), 404
    return jsonify({"message": "Eliminada"}), 200

# ARCHIVOS ADJUNTOS
@app.route("/api/archivos_adjuntos", methods=["GET"])
@require_auth
def get_archivos():
    alc_id = request.args.get("alcantarilla_id")
    query = supabase.table("archivos_adjuntos").select("*")
    if alc_id: query = query.eq("alcantarilla_id", alc_id)
    return jsonify(query.execute().data)

@app.route("/api/archivos_adjuntos/<int:id>", methods=["GET"])
@require_auth
def get_archivo(id):
    data = supabase.table("archivos_adjuntos").select("*").eq("id", id).execute()
    if not data.data: return jsonify({"detail": "No encontrado"}), 404
    return jsonify(data.data[0])

@app.route("/api/archivos_adjuntos", methods=["POST"])
@require_auth
def create_archivo():
    body = request.json or {}
    required = ["alcantarilla_id", "archivo_url"]
    if not all(k in body for k in required): return jsonify({"detail": f"Campos obligatorios: {required}"}), 400
    result = supabase.table("archivos_adjuntos").insert(body).execute()
    return jsonify(result.data[0]), 201

@app.route("/api/archivos_adjuntos/<int:id>", methods=["DELETE"])
@require_auth
def delete_archivo(id):
    result = supabase.table("archivos_adjuntos").delete().eq("id", id).execute()
    if not result.data: return jsonify({"detail": "No encontrado"}), 404
    return jsonify({"message": "Eliminado"}), 200

# ═════════════════════════════════════════════
# DASHBOARD
# ═════════════════════════════════════════════
@app.route("/api/dashboard/kpis", methods=["GET"])
@require_auth
def dashboard_kpis():
    total = supabase.table("alcantarillas").select("id", count="exact").execute().count or 0
    estados = supabase.table("estados").select("*").execute().data
    por_estado = []
    for e in estados:
        alc_ids = set()
        for tabla in ["tuberias", "muros", "pozos_recoleccion"]:
            rows = supabase.table(tabla).select("alcantarilla_id").eq("estado_id", e["id"]).execute().data
            for r in rows:
                if r.get("alcantarilla_id"): alc_ids.add(r["alcantarilla_id"])
        por_estado.append({"estado": e["nombre"], "cantidad": len(alc_ids)})
    alcs = supabase.table("alcantarillas").select("provincia").execute().data
    prov_count = {}
    for a in alcs:
        p = a.get("provincia") or "Sin provincia"
        prov_count[p] = prov_count.get(p, 0) + 1
    por_provincia = sorted([{"provincia": k, "cantidad": v} for k, v in prov_count.items()], key=lambda x: -x["cantidad"])[:8]
    materiales = supabase.table("materiales").select("*").execute().data
    por_material = []
    for m in materiales:
        cnt = supabase.table("tuberias").select("id", count="exact").eq("material_id", m["id"]).execute().count or 0
        if cnt > 0: por_material.append({"material": m["nombre"], "cantidad": cnt})
    ultimas = supabase.table("inspecciones").select("*").order("fecha", desc=True).limit(5).execute().data
    for ins in ultimas:
        alc = supabase.table("alcantarillas").select("ficha_numero, ubicacion").eq("id", ins["alcantarilla_id"]).execute().data
        ins["ficha_numero"] = alc[0]["ficha_numero"] if alc else None
        ins["ubicacion"] = alc[0]["ubicacion"] if alc else None
    return jsonify({
        "total_alcantarillas": total,
        "por_estado": por_estado,
        "por_provincia": por_provincia,
        "por_material": por_material,
        "ultimas_inspecciones": ultimas,
    })

# ═════════════════════════════════════════════
# MAPA GIS
# ═════════════════════════════════════════════
@app.route("/api/mapa/puntos", methods=["GET"])
@require_auth
def mapa_puntos():
    provincia = request.args.get("provincia")
    query = supabase.table("alcantarillas").select(
        "id, ficha_numero, ubicacion, canton, provincia, tramo_vial, coordenada_este, coordenada_norte"
    )
    if provincia:
        query = query.ilike("provincia", f"%{provincia}%")
    query = query.not_.is_("coordenada_este", "null").not_.is_("coordenada_norte", "null")
    rows = query.limit(2000).execute().data
    features = []
    for r in rows:
        este = r.get("coordenada_este")
        norte = r.get("coordenada_norte")
        try:
            este = float(este)
            norte = float(norte)
        except (TypeError, ValueError):
            continue
        features.append({
            "type": "Feature",
            "geometry": {
                "type": "Point",
                "coordinates": [este, norte]
            },
            "properties": {
                "id": r.get("id"),
                "ficha": r.get("ficha_numero"),
                "ubicacion": r.get("ubicacion"),
                "canton": r.get("canton"),
                "provincia": r.get("provincia"),
                "tramo_vial": r.get("tramo_vial"),
                "projection": "utm"
            }
        })
    return jsonify({
        "type": "FeatureCollection",
        "features": features,
        "total": len(features)
    })

# ═════════════════════════════════════════════
# FUNCIÓN AUXILIAR PARA EXPORTACIONES COMPLETAS
# ═════════════════════════════════════════════
def _get_full_alcantarillas():
    """Obtiene alcantarillas con tuberías, muros y pozos relacionados."""
    rows = supabase.table("alcantarillas").select("*").order("ficha_numero").execute().data
    full = []
    for r in rows:
        aid = r["id"]
        r["tuberias"] = supabase.table("tuberias").select(
            "*, materiales(nombre), estados(nombre)"
        ).eq("alcantarilla_id", aid).execute().data
        r["muros"] = supabase.table("muros").select(
            "*, materiales(nombre), estados(nombre)"
        ).eq("alcantarilla_id", aid).execute().data
        r["pozos"] = supabase.table("pozos_recoleccion").select(
            "*, estados(nombre)"
        ).eq("alcantarilla_id", aid).execute().data
        full.append(r)
    return full

# ═════════════════════════════════════════════
# EXCEL – HOJA RESUMEN + FICHAS DETALLADAS
# ═════════════════════════════════════════════
@app.route("/api/export/alcantarillas/excel", methods=["GET"])
@require_auth
def export_excel():
    rows = _get_full_alcantarillas()
    wb = Workbook()

    AZUL_OSC  = "1F3864"
    AZUL_MED  = "2E75B6"
    AZUL_CLAR = "BDD7EE"
    GRIS      = "D6DCE4"
    BLANCO    = "FFFFFF"
    NARANJA   = "ED7D31"

    def make_border():
        s = Side(style='thin', color='000000')
        return Border(left=s, right=s, top=s, bottom=s)

    def hdr_cell(ws, row, col, value, bg=AZUL_OSC, fg=BLANCO, bold=True, size=9, align='center'):
        c = ws.cell(row=row, column=col, value=value)
        c.font = Font(name='Calibri', bold=bold, color=fg, size=size)
        c.fill = PatternFill('solid', fgColor=bg)
        c.alignment = Alignment(horizontal=align, vertical='center')
        c.border = make_border()
        return c

    def data_cell(ws, row, col, value, bg=BLANCO, size=8, align='center'):
        c = ws.cell(row=row, column=col, value=value)
        c.font = Font(name='Calibri', color='000000', size=size)
        c.fill = PatternFill('solid', fgColor=bg)
        c.alignment = Alignment(horizontal=align, vertical='center')
        c.border = make_border()
        return c

    # ─── Hoja 1: Resumen ──────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Resumen'
    ws.sheet_view.showGridLines = False

    ws.merge_cells('A1:N1')
    ws['A1'].value = 'PONTIFICIA UNIVERSIDAD CATÓLICA DEL ECUADOR'
    ws['A1'].font = Font(name='Calibri', bold=True, color=BLANCO, size=14)
    ws['A1'].fill = PatternFill('solid', fgColor=AZUL_OSC)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    ws.merge_cells('A2:N2')
    ws['A2'].value = 'INVENTARIO DE ALCANTARILLAS — FICHA TÉCNICA RESUMEN'
    ws['A2'].font = Font(name='Calibri', bold=True, color=BLANCO, size=11)
    ws['A2'].fill = PatternFill('solid', fgColor=AZUL_MED)
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20

    headers = [
        ("N° Ficha", 6), ("Universidad", 18), ("Ubicación", 16),
        ("Parroquia", 12), ("Cantón", 12), ("Provincia", 12),
        ("Fecha", 9), ("Tramo Vial", 20),
        ("Mat. Tubería", 12), ("Long. Tub.(m)", 10), ("Diám. Tub.(m)", 10),
        ("Est. Tubería", 10), ("Pozo Recol.", 9), ("Coordenadas UTM", 18)
    ]
    for i, (label, width) in enumerate(headers, 1):
        hdr_cell(ws, 3, i, label, bg=AZUL_MED, size=8)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[3].height = 30

    for ri, r in enumerate(rows):
        row_num = ri + 4
        bg = AZUL_CLAR if ri % 2 == 0 else GRIS

        tub  = r["tuberias"][0]  if r["tuberias"]  else {}
        pozo = r["pozos"][0]     if r["pozos"]      else {}

        mat_tub  = (tub.get("materiales") or {}).get("nombre", "—")
        est_tub  = (tub.get("estados")    or {}).get("nombre", "—")
        long_tub = tub.get("longitud", "—")
        diam_tub = tub.get("diametro", "—")
        pozo_si  = "Sí" if pozo.get("existe") else ("No" if pozo else "—")
        coord    = f"E:{r.get('coordenada_este','')} / N:{r.get('coordenada_norte','')}" if r.get('coordenada_este') else "—"

        vals = [
            r.get("ficha_numero") or '—',
            r.get("universidad") or '—',
            r.get("ubicacion") or '—',
            r.get("parroquia") or '—',
            r.get("canton") or '—',
            r.get("provincia") or '—',
            r.get("fecha") or '—',
            r.get("tramo_vial") or '—',
            mat_tub, long_tub, diam_tub, est_tub, pozo_si, coord
        ]
        for ci, v in enumerate(vals, 1):
            align = 'left' if ci in (2,3,4,5,6,8) else 'center'
            data_cell(ws, row_num, ci, v, bg=bg, align=align)
        ws.row_dimensions[row_num].height = 16

    total_row = len(rows) + 4
    ws.merge_cells(f'A{total_row}:G{total_row}')
    ws[f'A{total_row}'].value = f'TOTAL DE REGISTROS: {len(rows)}'
    ws[f'A{total_row}'].font = Font(name='Calibri', bold=True, color=BLANCO, size=9)
    ws[f'A{total_row}'].fill = PatternFill('solid', fgColor=AZUL_OSC)
    ws[f'A{total_row}'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[total_row].height = 18

    # ─── Hoja 2: Fichas Detalladas ────────────────────────────────────
    ws2 = wb.create_sheet('Fichas Detalladas')
    ws2.sheet_view.showGridLines = False
    for i, w in enumerate([14]*8, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    current_row = 1
    for r in rows:
        tub  = r["tuberias"][0]  if r["tuberias"]  else {}
        muro = r["muros"][0]     if r["muros"]     else {}
        pozo = r["pozos"][0]     if r["pozos"]      else {}

        CR = current_row

        # Título
        ws2.merge_cells(f'A{CR}:F{CR}')
        hdr_cell(ws2, CR, 1, 'FICHA TÉCNICA - ALCANTARILLA', bg=AZUL_MED, size=11)
        ws2.merge_cells(f'G{CR}:H{CR}')
        hdr_cell(ws2, CR, 7, f"N° {r.get('ficha_numero') or '—'}", bg=AZUL_MED, size=11)
        ws2.row_dimensions[CR].height = 22
        CR += 1

        # Ubicación, Parroquia, Cantón, Provincia
        for ci, (lbl, val) in enumerate([
            ("Ubicación", r.get("ubicacion") or '—'),
            ("Parroquia", r.get("parroquia") or '—'),
            ("Cantón",    r.get("canton") or '—'),
            ("Provincia", r.get("provincia") or '—')
        ], 1):
            col = (ci-1)*2 + 1
            hdr_cell(ws2, CR, col, lbl, bg=AZUL_MED, size=8)
            data_cell(ws2, CR+1, col, val, align='left')
        ws2.row_dimensions[CR].height = 14
        ws2.row_dimensions[CR+1].height = 14
        CR += 2

        # Fecha, Tramo Vial
        hdr_cell(ws2, CR, 1, 'Fecha', bg=AZUL_MED, size=8)
        data_cell(ws2, CR, 2, r.get('fecha') or '—')
        hdr_cell(ws2, CR, 3, 'Tramo Vial', bg=AZUL_MED, size=8)
        ws2.merge_cells(f'D{CR}:H{CR}')
        data_cell(ws2, CR, 4, r.get('tramo_vial') or '—', align='left')
        ws2.row_dimensions[CR].height = 14
        CR += 1

        # INFRAESTRUCTURA EXISTENTE
        ws2.merge_cells(f'A{CR}:H{CR}')
        hdr_cell(ws2, CR, 1, 'INFRAESTRUCTURA EXISTENTE', bg=AZUL_OSC, size=9)
        ws2.row_dimensions[CR].height = 16
        CR += 1

        # MUROS DE ALA / TUBERÍA
        ws2.merge_cells(f'A{CR}:D{CR}')
        hdr_cell(ws2, CR, 1, 'MUROS DE ALA', bg=AZUL_MED, size=9)
        ws2.merge_cells(f'E{CR}:H{CR}')
        hdr_cell(ws2, CR, 5, 'TUBERÍA', bg=AZUL_MED, size=9)
        ws2.row_dimensions[CR].height = 16
        CR += 1

        ws2.merge_cells(f'A{CR}:B{CR}')
        hdr_cell(ws2, CR, 1, 'Dimensiones', bg=AZUL_MED, size=8)
        ws2.merge_cells(f'C{CR}:D{CR}')
        hdr_cell(ws2, CR, 3, 'Material', bg=AZUL_MED, size=8)
        # Eliminado merge E:H que causaba el error
        hdr_cell(ws2, CR, 5, 'Material', bg=AZUL_MED, size=8)  # Ahora solo escribe en E
        ws2.row_dimensions[CR].height = 14
        CR += 1

        mat_muro = (muro.get("materiales") or {}).get("nombre","—")
        mat_tub  = (tub.get("materiales")  or {}).get("nombre","—")

        hdr_cell(ws2, CR, 1, 'Longitud', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 2, f"{muro.get('longitud','NA')} m")
        hdr_cell(ws2, CR, 3, mat_muro, bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 4, '')
        # Ahora escribimos Cemento en E, y el ☑/☐ en F sin merge
        hdr_cell(ws2, CR, 5, 'Cemento', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 6, '☑' if mat_tub=='Cemento' else '☐')
        hdr_cell(ws2, CR, 7, 'PVC', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 8, '☑' if mat_tub=='PVC' else '☐')
        ws2.row_dimensions[CR].height = 14
        CR += 1

        hdr_cell(ws2, CR, 1, 'Espesor', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 2, f"{muro.get('espesor','NA')} m")
        ws2.merge_cells(f'E{CR}:H{CR}')
        hdr_cell(ws2, CR, 5, 'Dimensiones', bg=AZUL_MED, size=8)
        ws2.row_dimensions[CR].height = 14
        CR += 1

        hdr_cell(ws2, CR, 1, 'Solera', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 2, '')
        hdr_cell(ws2, CR, 3, 'H.Armado', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 4, '☑' if mat_muro=='Hormigón Armado' else '☐')
        hdr_cell(ws2, CR, 5, 'Longitud', bg=GRIS, fg='000000', bold=False, size=8)
        ws2.merge_cells(f'F{CR}:G{CR}')
        data_cell(ws2, CR, 6, f"{tub.get('longitud','—')} m")
        ws2.row_dimensions[CR].height = 14
        CR += 1

        hdr_cell(ws2, CR, 1, 'Muro gavión', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 2, '')
        hdr_cell(ws2, CR, 3, 'Estado', bg=GRIS, fg='000000', bold=False, size=8)
        est_muro = (muro.get("estados") or {}).get("nombre","—")
        data_cell(ws2, CR, 4, est_muro)
        hdr_cell(ws2, CR, 5, 'Diámetro', bg=GRIS, fg='000000', bold=False, size=8)
        ws2.merge_cells(f'F{CR}:G{CR}')
        data_cell(ws2, CR, 6, f"{tub.get('diametro','—')} m")
        ws2.row_dimensions[CR].height = 14
        CR += 1

        # MURO CABEZAL
        ws2.merge_cells(f'A{CR}:H{CR}')
        hdr_cell(ws2, CR, 1, 'MURO CABEZAL', bg=AZUL_MED, size=9)
        ws2.row_dimensions[CR].height = 16
        CR += 1

        for lbl_d, key_d in [("Longitud", "longitud"), ("Espesor", "espesor")]:
            hdr_cell(ws2, CR, 1, 'Dimensiones', bg=AZUL_MED, size=8)
            hdr_cell(ws2, CR, 2, lbl_d, bg=GRIS, fg='000000', bold=False, size=8)
            data_cell(ws2, CR, 3, muro.get(key_d, 'NA'))
            if lbl_d == "Longitud":
                hdr_cell(ws2, CR, 4, 'Estado', bg=AZUL_MED, size=8)
                for ci2, lbl2 in enumerate(["Bueno","Regular","Malo"], 5):
                    hdr_cell(ws2, CR, ci2, lbl2, bg=GRIS, fg='000000', bold=False, size=8)
                    est_m = (muro.get("estados") or {}).get("nombre","")
                    data_cell(ws2, CR, ci2, '☑' if est_m==lbl2 else '☐')
            else:
                data_cell(ws2, CR, 4, '')
            ws2.row_dimensions[CR].height = 14
            CR += 1

        # POZO DE RECOLECCIÓN
        ws2.merge_cells(f'A{CR}:H{CR}')
        hdr_cell(ws2, CR, 1, 'POZO DE RECOLECCIÓN', bg=AZUL_MED, size=9)
        ws2.row_dimensions[CR].height = 16
        CR += 1

        tiene_pozo = pozo.get("existe", False) if pozo else False
        hdr_cell(ws2, CR, 1, 'Sí', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 2, '☑' if tiene_pozo else '☐')
        hdr_cell(ws2, CR, 3, 'Dimensiones', bg=AZUL_MED, size=8)
        hdr_cell(ws2, CR, 4, 'Ancho', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 5, pozo.get('ancho','NA') if pozo else 'NA')
        hdr_cell(ws2, CR, 6, 'ESTADO', bg=AZUL_MED, size=8)
        est_p = (pozo.get("estados") or {}).get("nombre","") if pozo else ""
        for ci2, lbl2 in enumerate(["Bueno","Regular","Malo"], 7):
            hdr_cell(ws2, CR, ci2, lbl2, bg=GRIS, fg='000000', bold=False, size=8)
        ws2.row_dimensions[CR].height = 14
        CR += 1

        hdr_cell(ws2, CR, 1, 'No', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 2, '☑' if not tiene_pozo else '☐')
        data_cell(ws2, CR, 3, '')
        hdr_cell(ws2, CR, 4, 'Largo', bg=GRIS, fg='000000', bold=False, size=8)
        data_cell(ws2, CR, 5, pozo.get('largo','NA') if pozo else 'NA')
        data_cell(ws2, CR, 6, '')
        for ci2, lbl2 in enumerate(["Bueno","Regular","Malo"], 7):
            data_cell(ws2, CR, ci2, '☑' if est_p==lbl2 else '☐')
        ws2.row_dimensions[CR].height = 14
        CR += 1

        # COORDENADAS UTM
        ws2.merge_cells(f'A{CR}:H{CR}')
        hdr_cell(ws2, CR, 1, 'COORDENADAS UTM', bg=AZUL_OSC, size=9)
        ws2.row_dimensions[CR].height = 14
        CR += 1

        hdr_cell(ws2, CR, 1, 'ESTE', bg=NARANJA, size=8)
        ws2.merge_cells(f'B{CR}:D{CR}')
        data_cell(ws2, CR, 2, r.get('coordenada_este',''))
        hdr_cell(ws2, CR, 5, 'NORTE', bg=NARANJA, size=8)
        ws2.merge_cells(f'F{CR}:H{CR}')
        data_cell(ws2, CR, 6, r.get('coordenada_norte',''))
        ws2.row_dimensions[CR].height = 14
        CR += 1

        # OBSERVACIONES
        hdr_cell(ws2, CR, 1, 'Observaciones', bg=AZUL_MED, size=8)
        ws2.merge_cells(f'B{CR}:H{CR}')
        data_cell(ws2, CR, 2, r.get('observaciones',''), align='left', size=8)
        ws2.row_dimensions[CR].height = 20
        CR += 1

        # Separador
        ws2.row_dimensions[CR].height = 8
        CR += 1
        current_row = CR

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=alcantarillas.xlsx"}
    )

# ═════════════════════════════════════════════
# PDF – FORMATO DETALLADO
# ═════════════════════════════════════════════
@app.route("/api/export/alcantarillas/pdf", methods=["GET"])
@require_auth
def export_pdf():
    rows = _get_full_alcantarillas()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            rightMargin=1.2*cm, leftMargin=1.2*cm,
                            topMargin=1.2*cm, bottomMargin=1.2*cm)

    AZUL_OSC  = colors.HexColor("#1F3864")
    AZUL_MED  = colors.HexColor("#2E75B6")
    AZUL_CLAR = colors.HexColor("#BDD7EE")
    GRIS_CLAR = colors.HexColor("#D6DCE4")
    GRIS_MED  = colors.HexColor("#F2F2F2")
    NARANJA   = colors.HexColor("#ED7D31")
    BLANCO    = colors.white
    NEGRO     = colors.black

    styles = getSampleStyleSheet()
    st_titulo = ParagraphStyle('titulo', fontSize=10, fontName='Helvetica-Bold',
                               textColor=BLANCO, alignment=TA_CENTER, leading=13)
    st_sub    = ParagraphStyle('sub', fontSize=8, fontName='Helvetica-Bold',
                               textColor=BLANCO, alignment=TA_CENTER, leading=11)
    st_label  = ParagraphStyle('label', fontSize=7, fontName='Helvetica-Bold',
                               textColor=NEGRO, alignment=TA_CENTER, leading=9)
    st_labelL = ParagraphStyle('labelL', fontSize=7, fontName='Helvetica-Bold',
                               textColor=NEGRO, alignment=TA_LEFT, leading=9)
    st_val    = ParagraphStyle('val', fontSize=7, fontName='Helvetica',
                               textColor=NEGRO, alignment=TA_LEFT, leading=9)
    st_valc   = ParagraphStyle('valc', fontSize=7, fontName='Helvetica',
                               textColor=NEGRO, alignment=TA_CENTER, leading=9)

    PW = A4[0] - 2.4*cm

    def P(text, style=None):
        return Paragraph(str(text) if text is not None else '—', style or st_val)
    def PL(text): return P(text, st_label)
    def PC(text): return P(text, st_valc)
    def PLL(text): return P(text, st_labelL)
    def chk(cond): return '☑' if cond else '☐'

    def thin_border():
        return [
            ('GRID', (0,0), (-1,-1), 0.5, NEGRO),
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,0), (-1,-1), 7),
            ('TOPPADDING', (0,0), (-1,-1), 2),
            ('BOTTOMPADDING', (0,0), (-1,-1), 2),
            ('LEFTPADDING', (0,0), (-1,-1), 3),
            ('RIGHTPADDING', (0,0), (-1,-1), 3),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]

    elements = []

    # Portada resumen
    hdr = Table(
        [[P('PONTIFICIA UNIVERSIDAD CATÓLICA DEL ECUADOR', st_titulo),
          P('FICHA TÉCNICA — RESUMEN DE ALCANTARILLAS', st_sub)]],
        colWidths=[PW*0.65, PW*0.35]
    )
    hdr.setStyle(TableStyle([
        ('BACKGROUND', (0,0),(0,0), AZUL_OSC),
        ('BACKGROUND', (1,0),(1,0), AZUL_MED),
        ('TOPPADDING', (0,0),(-1,-1), 6),
        ('BOTTOMPADDING', (0,0),(-1,-1), 6),
        ('GRID', (0,0),(-1,-1), 0.5, NEGRO),
    ]))
    elements.append(hdr)
    elements.append(Spacer(1, 4))

    col_w = [PW*0.05, PW*0.08, PW*0.14, PW*0.10, PW*0.10, PW*0.10,
             PW*0.08, PW*0.15, PW*0.10, PW*0.10]
    resumen_data = [[
        PL('N°'), PL('Ficha'), PL('Ubicación'), PL('Parroquia'), PL('Cantón'),
        PL('Provincia'), PL('Fecha'), PL('Tramo Vial'), PL('Mat.Tub.'), PL('Est.Tub.')
    ]]
    for idx, r in enumerate(rows):
        tub = r['tuberias'][0] if r['tuberias'] else {}
        mat = (tub.get('materiales') or {}).get('nombre','—')
        est = (tub.get('estados') or {}).get('nombre','—')
        bg = AZUL_CLAR if idx % 2 == 0 else GRIS_CLAR
        resumen_data.append([
            PC(str(idx+1)),
            PC(r.get('ficha_numero') or '—'),
            P(r.get('ubicacion') or '—'),
            P(r.get('parroquia') or '—'),
            P(r.get('canton') or '—'),
            P(r.get('provincia') or '—'),
            PC(str(r.get('fecha') or '—')),
            P(r.get('tramo_vial') or '—'),
            PC(mat), PC(est),
        ])

    t_res = Table(resumen_data, colWidths=col_w, repeatRows=1)
    cmds = thin_border() + [
        ('BACKGROUND', (0,0),(-1,0), AZUL_MED),
        ('TEXTCOLOR', (0,0),(-1,0), BLANCO),
        ('FONTNAME', (0,0),(-1,0), 'Helvetica-Bold'),
    ]
    for i in range(1, len(resumen_data)):
        bg = AZUL_CLAR if i % 2 == 1 else GRIS_CLAR
        cmds.append(('BACKGROUND', (0,i),(-1,i), bg))
    t_res.setStyle(TableStyle(cmds))
    elements.append(t_res)

    total_row = Table(
        [[P(f'TOTAL DE REGISTROS: {len(rows)}', st_sub)]],
        colWidths=[PW]
    )
    total_row.setStyle(TableStyle([
        ('BACKGROUND', (0,0),(-1,-1), AZUL_OSC),
        ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),
    ]))
    elements.append(total_row)
    elements.append(PageBreak())

    # Fichas individuales
    for r in rows:
        tub  = r['tuberias'][0]  if r['tuberias']  else {}
        muro = r['muros'][0]     if r['muros']     else {}
        pozo = r['pozos'][0]     if r['pozos']      else {}
        mat_tub  = (tub.get('materiales') or {}).get('nombre','—')
        est_tub  = (tub.get('estados') or {}).get('nombre','—')
        mat_muro = (muro.get('materiales') or {}).get('nombre','—')
        est_muro = (muro.get('estados') or {}).get('nombre','—')
        est_pozo = (pozo.get('estados') or {}).get('nombre','') if pozo else ''
        tiene_p  = pozo.get('existe', False) if pozo else False

        ficha = []

        # Encabezado
        enc = Table([
            [P('PONTIFICIA UNIVERSIDAD CATÓLICA DEL ECUADOR', st_titulo),
             P(f"N° {r.get('ficha_numero') or '—'}", st_sub)]
        ], colWidths=[PW*0.75, PW*0.25])
        enc.setStyle(TableStyle([
            ('BACKGROUND', (0,0),(0,0), AZUL_OSC),
            ('BACKGROUND', (1,0),(1,0), AZUL_MED),
            ('GRID', (0,0),(-1,-1), 0.5, NEGRO),
            ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),
        ]))
        ficha.append(enc)

        ft = Table([[P('FICHA TÉCNICA - ALCANTARILLA', st_sub)]], colWidths=[PW])
        ft.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,-1), AZUL_MED),
            ('GRID',(0,0),(-1,-1),0.5,NEGRO),
            ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),
        ]))
        ficha.append(ft)

        # Datos de ubicación
        c = PW / 8
        t_ub = Table([
            [PL('Ubicación'), PC(r.get('ubicacion') or '—'),
             PL('Parroquia'), PC(r.get('parroquia') or '—'),
             PL('Cantón'),    PC(r.get('canton') or '—'),
             PL('Provincia'), PC(r.get('provincia') or '—')]
        ], colWidths=[c*0.8, c*1.2, c*0.8, c*1.2, c*0.7, c*1.1, c*0.8, c*1.4])
        t_ub.setStyle(TableStyle(thin_border() + [
            ('BACKGROUND',(0,0),(0,0),AZUL_MED),('TEXTCOLOR',(0,0),(0,0),BLANCO),
            ('BACKGROUND',(2,0),(2,0),AZUL_MED),('TEXTCOLOR',(2,0),(2,0),BLANCO),
            ('BACKGROUND',(4,0),(4,0),AZUL_MED),('TEXTCOLOR',(4,0),(4,0),BLANCO),
            ('BACKGROUND',(6,0),(6,0),AZUL_MED),('TEXTCOLOR',(6,0),(6,0),BLANCO),
        ]))
        ficha.append(t_ub)

        t_ft2 = Table([
            [PL('Fecha'), PC(str(r.get('fecha') or '—')),
             PL('Tramo vial'), P(r.get('tramo_vial') or '—')]
        ], colWidths=[PW*0.08, PW*0.17, PW*0.1, PW*0.65])
        t_ft2.setStyle(TableStyle(thin_border() + [
            ('BACKGROUND',(0,0),(0,0),AZUL_MED),('TEXTCOLOR',(0,0),(0,0),BLANCO),
            ('BACKGROUND',(2,0),(2,0),AZUL_MED),('TEXTCOLOR',(2,0),(2,0),BLANCO),
        ]))
        ficha.append(t_ft2)

        inf = Table([[P('INFRAESTRUCTURA EXISTENTE', st_sub)]], colWidths=[PW])
        inf.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,-1),AZUL_OSC),
            ('GRID',(0,0),(-1,-1),0.5,NEGRO),
            ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),
        ]))
        ficha.append(inf)

        # Muros de Ala y Tubería
        LW = PW * 0.5 - 0.5
        RW = PW * 0.5 - 0.5
        lc = LW / 4
        muro_rows = [
            [PL('MUROS DE ALA'), P(''), PL('TUBERÍA'), P('')],
            [PL('Dimensiones'), P(''), PL('Material'), P('')],
            [PLL('Longitud'), PC(f"{muro.get('longitud','NA')} m"),
             PLL('Cemento'), PC(chk(mat_tub=='Cemento'))],
            [PLL('Espesor'),  PC(f"{muro.get('espesor','NA')} m"),
             PLL('PVC'),      PC(chk(mat_tub=='PVC'))],
            [PLL('Material muro'), PC(mat_muro),
             PLL('M.Corrugado'),   PC(chk(mat_tub=='Metal Corrugado'))],
            [PL('Dimensiones Tubería'), P(''), PL('Estado Tubería'), P('')],
            [PLL('Solera'),      PC(chk(False)),
             PLL('Longitud'),    PC(f"{tub.get('longitud','—')} m")],
            [PLL('H.Armado'),    PC(chk(mat_muro=='Hormigón Armado')),
             PLL('Diámetro'),    PC(f"{tub.get('diametro','—')} m")],
            [PLL('Muro gavión'), PC(chk(False)),
             PLL('Estado'),      PC(est_tub)],
        ]
        t_muros = Table(muro_rows, colWidths=[lc*1.3, lc*0.7, lc*1.3, lc*0.7])
        cmds_m = thin_border() + [
            ('SPAN',(0,0),(1,0)),('SPAN',(2,0),(3,0)),
            ('BACKGROUND',(0,0),(1,0),AZUL_MED),('TEXTCOLOR',(0,0),(1,0),BLANCO),
            ('BACKGROUND',(2,0),(3,0),AZUL_MED),('TEXTCOLOR',(2,0),(3,0),BLANCO),
            ('FONTNAME',(0,0),(3,0),'Helvetica-Bold'),
            ('SPAN',(0,1),(1,1)),('SPAN',(2,1),(3,1)),
            ('BACKGROUND',(0,1),(1,1),AZUL_CLAR),
            ('BACKGROUND',(2,1),(3,1),AZUL_CLAR),
            ('FONTNAME',(0,1),(3,1),'Helvetica-Bold'),
            ('SPAN',(0,5),(1,5)),('SPAN',(2,5),(3,5)),
            ('BACKGROUND',(0,5),(1,5),AZUL_CLAR),
            ('BACKGROUND',(2,5),(3,5),AZUL_CLAR),
            ('FONTNAME',(0,5),(3,5),'Helvetica-Bold'),
        ]
        t_muros.setStyle(TableStyle(cmds_m))

        # Muro Cabezal
        mc_rows = [
            [P('MURO CABEZAL', st_sub), P(''), P('')],
            [PL('Dimensiones'), PLL('Longitud'), PC('NA')],
            [P(''), PLL('Espesor'),  PC('NA')],
            [PL('Estado'), PC(chk(est_muro=='Bueno')+' Bueno'),
             PC(chk(est_muro=='Regular')+' Regular')],
            [P(''), PC(chk(est_muro=='Malo')+' Malo'), P('')],
        ]
        t_mc = Table(mc_rows, colWidths=[RW*0.35, RW*0.35, RW*0.30])
        t_mc.setStyle(TableStyle(thin_border() + [
            ('SPAN',(0,0),(-1,0)),
            ('BACKGROUND',(0,0),(-1,0),AZUL_MED),
            ('BACKGROUND',(0,1),(0,2),AZUL_CLAR),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ]))

        t_infra = Table([[t_muros, t_mc]], colWidths=[LW+0.5, RW+0.5])
        t_infra.setStyle(TableStyle([
            ('VALIGN',(0,0),(-1,-1),'TOP'),
            ('LEFTPADDING',(0,0),(-1,-1),0),
            ('RIGHTPADDING',(0,0),(-1,-1),0),
        ]))
        ficha.append(t_infra)

        # Estado
        estado_rows = [
            [PL('Estado'), PL('B'), PL('R'), PL('M'), PL('B'), PL('R'), PL('M')],
            [P(''),
             PC(chk(est_muro=='Bueno')), PC(chk(est_muro=='Regular')), PC(chk(est_muro=='Malo')),
             PC(chk(est_muro=='Bueno')), PC(chk(est_muro=='Regular')), PC(chk(est_muro=='Malo'))],
        ]
        cw_e = PW/7
        t_est = Table(estado_rows, colWidths=[cw_e]*7)
        t_est.setStyle(TableStyle(thin_border() + [
            ('BACKGROUND',(0,0),(0,1),AZUL_CLAR),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ]))
        ficha.append(t_est)

        # Pozo
        pz = Table([[P('Pozo de recoleccion', st_sub)]], colWidths=[PW])
        pz.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,-1),AZUL_MED),
            ('GRID',(0,0),(-1,-1),0.5,NEGRO),
            ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),
        ]))
        ficha.append(pz)

        cw_p = PW / 8
        pozo_rows = [
            [PLL('Si'),  PC(chk(tiene_p)),
             PL('Dimensiones'), PLL('Ancho'), PC(str(pozo.get('ancho','NA')) if pozo else 'NA'),
             PL('ESTADO'),
             PC(chk(est_pozo=='Bueno')+' Bueno'),
             PC(chk(est_pozo=='Regular')+' Reg.')],
            [PLL('NO'), PC(chk(not tiene_p)),
             P(''), PLL('Largo'), PC(str(pozo.get('largo','NA')) if pozo else 'NA'),
             P(''),
             PC(chk(est_pozo=='Malo')+' Malo'), P('')],
        ]
        t_pz = Table(pozo_rows, colWidths=[cw_p*0.5, cw_p*0.5, cw_p*1.2, cw_p*0.7, cw_p*0.9, cw_p*0.8, cw_p*1.7, cw_p*1.7])
        t_pz.setStyle(TableStyle(thin_border() + [
            ('BACKGROUND',(2,0),(2,1),AZUL_MED),('TEXTCOLOR',(2,0),(2,1),BLANCO),
            ('SPAN',(2,0),(2,1)),
            ('BACKGROUND',(5,0),(5,1),AZUL_MED),('TEXTCOLOR',(5,0),(5,1),BLANCO),
            ('SPAN',(5,0),(5,1)),
        ]))
        ficha.append(t_pz)

        # Fotografía placeholder
        foto_label = Table([[PL('Fotografía')]], colWidths=[PW])
        foto_label.setStyle(TableStyle([
            ('GRID',(0,0),(-1,-1),0.5,NEGRO),
            ('BACKGROUND',(0,0),(-1,-1),GRIS_MED),
        ]))
        ficha.append(foto_label)
        foto_box = Table([[PC('[Espacio para fotografía]')]], colWidths=[PW])
        foto_box.setStyle(TableStyle([
            ('GRID',(0,0),(-1,-1),0.5,NEGRO),
            ('ROWHEIGHT',(0,0),(0,0),80),
            ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ]))
        ficha.append(foto_box)

        # Coordenadas
        coord_rows = [
            [PL('ESTE'),  PC(str(r.get('coordenada_este','—'))),
             PC(str(r.get('coordenada_este','—'))),
             PL('NORTE'), PC(str(r.get('coordenada_norte','—'))),
             PC(str(r.get('coordenada_norte','—')))],
        ]
        t_coord = Table(
            [[P('Cordenadas UTM', st_labelL)] + [P('')]*5, coord_rows[0]],
            colWidths=[PW*0.15, PW*0.14, PW*0.14, PW*0.15, PW*0.21, PW*0.21]
        )
        t_coord.setStyle(TableStyle(thin_border() + [
            ('SPAN',(0,0),(-1,0)),
            ('BACKGROUND',(0,0),(-1,0),AZUL_CLAR),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
            ('BACKGROUND',(0,1),(0,1),NARANJA),
            ('BACKGROUND',(3,1),(3,1),NARANJA),
            ('FONTNAME',(0,1),(0,1),'Helvetica-Bold'),('TEXTCOLOR',(0,1),(0,1),BLANCO),
            ('FONTNAME',(3,1),(3,1),'Helvetica-Bold'),('TEXTCOLOR',(3,1),(3,1),BLANCO),
        ]))
        ficha.append(t_coord)

        # Observaciones
        obs = r.get('observaciones','') or ''
        t_obs = Table([[PLL('Observaciones'), P(obs)]], colWidths=[PW*0.15, PW*0.85])
        t_obs.setStyle(TableStyle(thin_border() + [
            ('BACKGROUND',(0,0),(0,0),AZUL_MED),('TEXTCOLOR',(0,0),(0,0),BLANCO),
            ('FONTNAME',(0,0),(0,0),'Helvetica-Bold'),
            ('ROWHEIGHT',(0,0),(0,0),24),
        ]))
        ficha.append(t_obs)

        elements.append(KeepTogether(ficha[:6]))
        for el in ficha[6:]:
            elements.append(el)
        elements.append(PageBreak())

    doc.build(elements)
    buffer.seek(0)

    return Response(
        buffer,
        mimetype="application/pdf",
        headers={"Content-Disposition": "attachment; filename=alcantarillas.pdf"}
    )

# ═════════════════════════════════════════════
# MAIN
# ═════════════════════════════════════════════
if __name__ == "__main__":
    port = int(os.getenv("PORT", 3000))
    debug = os.getenv("FLASK_DEBUG", "1") == "1"
    app.run(host="0.0.0.0", port=port, debug=debug)